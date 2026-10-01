from __future__ import annotations

import html
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer

from .clarity import audit_novice_document

NOTE_MARKER = "Nota scientifica aggiornata"


@dataclass(frozen=True, slots=True)
class ScientificNote:
    key: str
    trigger: re.Pattern[str]
    title: str
    text: str
    source_label: str
    source_url: str


@dataclass(frozen=True, slots=True)
class FinalizationResult:
    markdown_path: Path
    docx_path: Path | None
    pdf_path: Path | None
    scientific_notes: int
    organization_fixes: int
    pdf_pages: int


SCIENTIFIC_NOTES: tuple[ScientificNote, ...] = (
    ScientificNote(
        key="systemic-venous-return",
        trigger=re.compile(
            r"Il sangue povero di ossigeno arriva al cuore tramite le arterie del corpo",
            re.IGNORECASE,
        ),
        title="Ritorno venoso sistemico",
        text=(
            "Nel circolo sistemico il sangue deossigenato ritorna al cuore attraverso il sistema "
            "venoso, fino alle vene cave superiore e inferiore e quindi all'atrio destro. "
            "Mantieni la formulazione eCampus sopra se serve per riconoscere la fonte, ma non "
            "memorizzarla come descrizione anatomica corretta."
        ),
        source_label="NCBI Bookshelf — Anatomy, Blood Flow",
        source_url="https://www.ncbi.nlm.nih.gov/sites/books/NBK554457/",
    ),
    ScientificNote(
        key="passive-transport-gradient",
        trigger=re.compile(
            r"meccanismi passivi.*?dirett[oa].*?zona a maggiore concentrazione",
            re.IGNORECASE | re.DOTALL,
        ),
        title="Direzione del trasporto passivo",
        text=(
            "Il flusso passivo netto procede lungo il gradiente: per una sostanza neutra da "
            "concentrazione maggiore verso concentrazione minore; per gli ioni conta il gradiente "
            "elettrochimico. Questa nota segnala una contraddizione presente nel testo sorgente."
        ),
        source_label="NCBI Bookshelf — Transport of Small Molecules",
        source_url="https://www.ncbi.nlm.nih.gov/books/NBK9847/",
    ),
    ScientificNote(
        key="organelles-not-biomolecules",
        trigger=re.compile(r"Gli organuli sono biomolecole", re.IGNORECASE),
        title="Organuli cellulari",
        text=(
            "Gli organuli non sono biomolecole: sono strutture o compartimenti intracellulari "
            "specializzati, come nucleo, reticolo endoplasmatico, Golgi, mitocondri e lisosomi."
        ),
        source_label="NCBI Bookshelf — The Compartmentalization of Cells",
        source_url="https://www.ncbi.nlm.nih.gov/books/NBK26907/",
    ),
    ScientificNote(
        key="parotid-endocrine",
        trigger=re.compile(
            r"Sistema endocrino \(composto da[^\n]*?parotidi",
            re.IGNORECASE,
        ),
        title="Parotidi",
        text=(
            "Le parotidi sono ghiandole salivari esocrine, non ghiandole endocrine. "
            "Non vanno confuse con le paratiroidi, che invece sono endocrine."
        ),
        source_label="NCBI Bookshelf — Anatomy, Head and Neck, Salivary Glands",
        source_url="https://www.ncbi.nlm.nih.gov/books/NBK538325/",
    ),
    ScientificNote(
        key="neurons-glia-count",
        trigger=re.compile(
            r"(?:100 milioni di neuroni|1 neurone\s*:\s*9 glie)",
            re.IGNORECASE,
        ),
        title="Numero di neuroni e rapporto glia/neuroni",
        text=(
            "Le stime moderne più citate indicano circa 86 miliardi di neuroni nel cervello umano "
            "adulto e un numero complessivo di cellule non neuronali dello stesso ordine di "
            "grandezza. I vecchi rapporti 9:1 o 10:1 non descrivono il cervello umano nel suo insieme."
        ),
        source_label="PubMed — Azevedo et al., 2009",
        source_url="https://pubmed.ncbi.nlm.nih.gov/19226510/",
    ),
    ScientificNote(
        key="lactate-fatigue",
        trigger=re.compile(
            r"accumulo di acido lattico.*?(?:bruciore|dolore).*?post-esercizio",
            re.IGNORECASE | re.DOTALL,
        ),
        title="Lattato, acidosi e fatica",
        text=(
            "La spiegazione tradizionale che attribuisce direttamente al lattato il bruciore, "
            "l'acidosi e la fatica è una semplificazione superata. La fatica è multifattoriale e "
            "la produzione di lattato non è, di per sé, la causa biochimica dell'acidosi."
        ),
        source_label="PubMed — Robergs et al., 2004",
        source_url="https://pubmed.ncbi.nlm.nih.gov/15308499/",
    ),
    ScientificNote(
        key="hemoglobin-lipoprotein",
        trigger=re.compile(
            r"lipoproteine \(proteine legate a lipidi\), come l['’]emoglobina",
            re.IGNORECASE,
        ),
        title="Emoglobina",
        text=(
            "L'emoglobina non è una lipoproteina: è una proteina tetramerica contenente gruppi "
            "eme con ferro, specializzata nel trasporto reversibile dell'ossigeno."
        ),
        source_label="NCBI Bookshelf — Oxygen Transport",
        source_url="https://www.ncbi.nlm.nih.gov/books/NBK54103/",
    ),
)


def _paragraph_ranges(lines: list[str]) -> Iterable[tuple[int, int, str]]:
    start: int | None = None
    for index, line in enumerate(lines + [""]):
        stripped = line.strip()
        if stripped and not stripped.startswith("#") and stripped != "---":
            if start is None:
                start = index
            continue
        if start is not None:
            end = index
            yield start, end, "\n".join(lines[start:end])
            start = None


def _apply_scientific_notes(text: str) -> tuple[str, int]:
    if NOTE_MARKER in text:
        return text, 0

    lines = text.splitlines()
    applied: set[str] = set()
    insertions: list[tuple[int, list[str]]] = []

    for _start, end, paragraph in _paragraph_ranges(lines):
        for note in SCIENTIFIC_NOTES:
            if note.key in applied or not note.trigger.search(paragraph):
                continue
            note_lines = [
                "",
                f"> **{NOTE_MARKER} — {note.title}.** {note.text}",
                f"> Fonte: {note.source_label} — {note.source_url}",
                "",
            ]
            insertions.append((end, note_lines))
            applied.add(note.key)

    for index, addition in sorted(insertions, reverse=True):
        lines[index:index] = addition

    return "\n".join(lines).rstrip() + "\n", len(applied)


def _chapter_slice(text: str, title: str) -> str:
    match = re.search(
        rf"(?ms)^##\s+{re.escape(title)}\s*$\n(?P<body>.*?)(?=^##\s+|\Z)",
        text,
    )
    return match.group("body") if match else ""


def _fix_known_mismatches(text: str) -> tuple[str, int]:
    pattern = re.compile(
        r"(?mi)^##\s+olfatto\s*$"
        r"(?P<warning>\n\s*>\s*Verifica materiale:[^\n]*"
        r"(?:\n\s*>[^\n]*)?)"
    )
    match = pattern.search(text)
    if not match:
        return text, 0

    gusto = _chapter_slice(text, "GUSTO")
    olfatto_body = _chapter_slice(text, "olfatto")
    if "controllo nervoso" not in olfatto_body.casefold():
        return text, 0

    replacement = (
        '## Controllo nervoso del movimento (materiale etichettato "olfatto")'
        + match.group("warning")
    )
    if "olfatt" in gusto.casefold():
        replacement += (
            "\n> **Nota organizzativa:** il contenuto propriamente olfattivo è già "
            "parzialmente trattato nel capitolo GUSTO; questo capitolo conserva invece il "
            "materiale di controllo motorio presente nel gruppo originariamente etichettato "
            '"olfatto".'
        )

    return text[: match.start()] + replacement + text[match.end() :], 1


def _finalize_frontmatter(text: str) -> tuple[str, int]:
    """Keep cover, map, recap, index and renamed chapter labels consistent."""
    changed = 0

    title_pattern = re.compile(
        r"(?m)^#\s+(.+?)\s*-\s*Riassunto rifinito\s*$"
    )
    if title_pattern.search(text):
        text = title_pattern.sub(r"# \1 - Dispensa finale", text, count=1)
        changed += 1

    old_subtitle = (
        "> Versione rifinita dell'ultimo riassunto BC Science; "
        "non sostituisce la verifica sui PDF originali."
    )
    new_subtitle = (
        "> Dispensa finale generata da BC Science. Il testo eCampus resta la base di studio; "
        "le eventuali precisazioni scientifiche sono separate e citate."
    )
    if old_subtitle in text:
        text = text.replace(old_subtitle, new_subtitle, 1)
        changed += 1

    old_label = "olfatto"
    new_label = 'Controllo nervoso del movimento (materiale etichettato "olfatto")'

    replacements = (
        (f"- {old_label}\n", f"- {new_label}\n"),
        (
            "gusto/tatto/olfatto/udito/vista",
            'gusto/tatto/udito/vista + controllo nervoso del movimento (file "olfatto")',
        ),
    )
    for old, new in replacements:
        if old in text:
            text = text.replace(old, new, 1)
            changed += 1

    recap_pattern = re.compile(
        r"(?m)^(?P<prefix>\d+\.\s+)\*\*olfatto(?:\s*\(titolo da verificare\))?\*\*:"
    )
    if recap_pattern.search(text):
        text = recap_pattern.sub(
            lambda match: f"{match.group('prefix')}**{new_label}**:",
            text,
            count=1,
        )
        changed += 1

    return text, changed


def _apply_scientific_notes_to_chapters(text: str) -> tuple[str, int]:
    """Attach scientific notes to full chapters, not to front-matter recaps."""
    index_pos = text.find("## Indice degli argomenti")
    if index_pos < 0:
        return _apply_scientific_notes(text)

    content_pos = text.find("\n---\n", index_pos)
    if content_pos < 0:
        return _apply_scientific_notes(text)

    split_at = content_pos + len("\n---\n")
    prefix = text[:split_at]
    chapters = text[split_at:]
    chapters, note_count = _apply_scientific_notes(chapters)
    return prefix + chapters, note_count


def finalize_markdown(text: str) -> tuple[str, int, int]:
    finalized, organization_fixes = _fix_known_mismatches(text)
    finalized, frontmatter_fixes = _finalize_frontmatter(finalized)
    organization_fixes += frontmatter_fixes
    finalized, scientific_notes = _apply_scientific_notes_to_chapters(finalized)
    if not finalized.endswith("\n"):
        finalized += "\n"
    return finalized, scientific_notes, organization_fixes


def _normalize_inline(text: str) -> str:
    replacements = {
        "$\\rightarrow$": "→",
        "$\\to$": "→",
        "\\rightarrow": "→",
        "\\ge": "≥",
        "\\le": "≤",
        "\\pm": "±",
    }
    value = text
    for old, new in replacements.items():
        value = value.replace(old, new)
    return value.replace("$", "")


def _plain_text(text: str) -> str:
    value = _normalize_inline(text)
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"(?<!\*)\*(.+?)\*(?!\*)", r"\1", value)
    return value


def _reportlab_inline(text: str) -> str:
    value = html.escape(_normalize_inline(text))
    value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"(?<!\*)\*(.+?)\*(?!\*)", r"<i>\1</i>", value)
    return value


def _register_pdf_fonts() -> tuple[str, str]:
    candidates = (
        (
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/arialbd.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/Library/Fonts/Arial.ttf"),
            Path("/Library/Fonts/Arial Bold.ttf"),
        ),
    )
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("BCScienceSans", str(regular)))
            pdfmetrics.registerFont(TTFont("BCScienceSansBold", str(bold)))
            return "BCScienceSans", "BCScienceSansBold"
    return "Helvetica", "Helvetica-Bold"


def export_pdf(markdown: str, destination: Path, title: str) -> int:
    regular_font, bold_font = _register_pdf_fonts()
    styles = getSampleStyleSheet()

    body = ParagraphStyle(
        "BCBody",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=10,
        leading=14,
        spaceAfter=5,
    )
    h1 = ParagraphStyle(
        "BCH1",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=23,
        leading=28,
        alignment=TA_CENTER,
        spaceAfter=14,
    )
    h2 = ParagraphStyle(
        "BCH2",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=16,
        leading=20,
        spaceBefore=5,
        spaceAfter=9,
    )
    h3 = ParagraphStyle(
        "BCH3",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=12,
        leading=15,
        spaceBefore=7,
        spaceAfter=5,
    )
    quote = ParagraphStyle(
        "BCQuote",
        parent=body,
        leftIndent=8 * mm,
        rightIndent=5 * mm,
        borderWidth=0.5,
        borderPadding=5,
        spaceBefore=4,
        spaceAfter=7,
    )
    note = ParagraphStyle(
        "BCNote",
        parent=quote,
        fontName=regular_font,
        fontSize=9,
        leading=12,
        borderWidth=0.8,
        borderPadding=6,
    )
    bullet = ParagraphStyle(
        "BCBullet",
        parent=body,
        leftIndent=7 * mm,
        firstLineIndent=-4 * mm,
    )

    destination.parent.mkdir(parents=True, exist_ok=True)

    def draw_page(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(regular_font, 8)
        canvas.drawString(18 * mm, 10 * mm, "BC Science · eCampus 2026/2027")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Pagina {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=17 * mm,
        title=title,
        author="BC Science",
    )

    story = []
    paragraph_buffer: list[str] = []
    seen_h1 = False
    seen_h2 = False

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        raw = " ".join(item.strip() for item in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if raw:
            story.append(Paragraph(_reportlab_inline(raw), body))

    for raw_line in markdown.splitlines():
        stripped = raw_line.strip()

        if not stripped:
            flush_paragraph()
            continue

        if stripped == "---":
            flush_paragraph()
            story.append(Spacer(1, 3))
            story.append(HRFlowable(width="100%", thickness=0.5))
            story.append(Spacer(1, 5))
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            if seen_h1:
                story.append(PageBreak())
            story.append(Paragraph(_reportlab_inline(stripped[2:]), h1))
            seen_h1 = True
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            if seen_h2:
                story.append(PageBreak())
            story.append(Paragraph(_reportlab_inline(stripped[3:]), h2))
            seen_h2 = True
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            story.append(Paragraph(_reportlab_inline(stripped[4:]), h3))
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_text = stripped.lstrip(">").strip()
            selected = note if NOTE_MARKER in quote_text else quote
            story.append(Paragraph(_reportlab_inline(quote_text), selected))
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            flush_paragraph()
            story.append(Paragraph("• " + _reportlab_inline(bullet_match.group(1)), bullet))
            continue

        numbered_match = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if numbered_match:
            flush_paragraph()
            story.append(
                Paragraph(
                    f"{numbered_match.group(1)}. "
                    + _reportlab_inline(numbered_match.group(2)),
                    bullet,
                )
            )
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph()
    doc.build(story, onFirstPage=draw_page, onLaterPages=draw_page)

    with pymupdf.open(destination) as pdf:
        if pdf.page_count <= 0:
            raise ValueError("Il PDF finale non contiene pagine.")
        return pdf.page_count


def export_docx(markdown: str, destination: Path, title: str) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    footer = section.footer.paragraphs[0]
    footer.text = "BC Science · eCampus 2026/2027"

    seen_h2 = False
    for raw_line in markdown.splitlines():
        stripped = raw_line.strip()
        if not stripped or stripped == "---":
            continue

        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.add_run(_plain_text(stripped[2:]))
            continue

        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            if seen_h2:
                paragraph.paragraph_format.page_break_before = True
            paragraph.add_run(_plain_text(stripped[3:]))
            seen_h2 = True
            continue

        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.add_run(_plain_text(stripped[4:]))
            continue

        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(_plain_text(stripped.lstrip(">").strip()))
            if NOTE_MARKER in stripped:
                run.bold = True
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet_match:
            document.add_paragraph(_plain_text(bullet_match.group(1)), style="List Bullet")
            continue

        numbered_match = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if numbered_match:
            document.add_paragraph(_plain_text(numbered_match.group(2)), style="List Number")
            continue

        document.add_paragraph(_plain_text(stripped))

    document.core_properties.title = title
    document.core_properties.author = "BC Science"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def finalize_file(
    source: Path,
    output_dir: Path,
    *,
    create_docx: bool = True,
    create_pdf: bool = True,
) -> FinalizationResult:
    source = source.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    if not source.exists() or source.suffix.lower() != ".md":
        raise ValueError(f"Riassunto Markdown non valido: {source}")

    raw = source.read_text(encoding="utf-8")
    novice_ready, novice_results = audit_novice_document(raw)
    if not novice_ready:
        failed = [
            f"{title}: {', '.join(audit.issues)}"
            for title, audit in novice_results
            if not audit.passed
        ]
        preview = "; ".join(failed[:5])
        more = f" (+{len(failed) - 5} altri)" if len(failed) > 5 else ""
        raise ValueError(
            "Il riassunto non supera il controllo di comprensibilita per principianti. "
            "Esegui prima 'bc-science refine' con la versione corrente. "
            f"Problemi: {preview}{more}"
        )

    final_markdown, note_count, organization_fixes = finalize_markdown(raw)

    stem = source.stem
    for suffix in ("-riassunto-rifinito", "-riassunto-unico", "-dispensa-finale"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break

    title = f"{stem} - Dispensa finale"
    output_dir.mkdir(parents=True, exist_ok=True)
    md_path = output_dir / f"{stem}-dispensa-finale.md"
    docx_path = output_dir / f"{stem}-dispensa-finale.docx" if create_docx else None
    pdf_path = output_dir / f"{stem}-dispensa-finale.pdf" if create_pdf else None

    md_path.write_text(final_markdown, encoding="utf-8")
    if docx_path:
        export_docx(final_markdown, docx_path, title)
    pages = export_pdf(final_markdown, pdf_path, title) if pdf_path else 0

    return FinalizationResult(
        markdown_path=md_path,
        docx_path=docx_path,
        pdf_path=pdf_path,
        scientific_notes=note_count,
        organization_fixes=organization_fixes,
        pdf_pages=pages,
    )
