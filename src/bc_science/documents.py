from __future__ import annotations

import base64
import hashlib
import os
import re
import shutil
import subprocess
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import pymupdf
import pymupdf4llm
from docx import Document

from .cache import CacheDB
from .config import AppConfig, app_home

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}
PDF_RECOVERY_TEXT_THRESHOLD = 80
DEFAULT_OCR_DPI = 150

GENERATED_STEM_SUFFIXES = (
    "-dispensa-finale",
    "-riassunto-unico",
    "-riassunto-studio",
)
GENERATED_STEMS = {
    "riassunto-unico",
    "riassunto-studio",
    "approfondimento-scientifico",
    "audit",
}
GENERATED_STEM_PREFIXES = ("benchmark-",)


@dataclass(slots=True)
class ExtractedDocument:
    path: Path
    text: str
    pages: int = 0
    ocr_pages: int = 0


def is_generated_artifact(path: Path) -> bool:
    """Return True for files produced by BC Science rather than primary study sources."""
    stem = path.stem.casefold()
    return (
        stem in GENERATED_STEMS
        or any(stem.endswith(suffix) for suffix in GENERATED_STEM_SUFFIXES)
        or any(stem.startswith(prefix) for prefix in GENERATED_STEM_PREFIXES)
    )


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_text(text: str) -> str:
    lines = [
        line.rstrip()
        for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]
    out: list[str] = []
    blank = False
    for line in lines:
        stripped = " ".join(line.split())
        if not stripped:
            if out and not blank:
                out.append("")
            blank = True
            continue
        out.append(stripped)
        blank = False
    return "\n".join(out).strip()


def _needs_pdf_recovery(text: str) -> bool:
    # Repeated course metadata can exceed the threshold while the actual
    # educational content is an image. Measure the substantive body instead.
    stripped = re.sub(
        r"(?mi)^\s*(?:Corso di Laurea|Insegnamento|Lezione n[°º.]?|Numero lezione|Titolo|Attivit[aà] n[°º.]?)\s*:[^\n]*$",
        "", text,
    ).strip()
    if len(stripped) < PDF_RECOVERY_TEXT_THRESHOLD:
        return True
    replacement_chars = stripped.count("\ufffd")
    return replacement_chars >= 3 and replacement_chars / len(stripped) >= 0.01


def _ocr_dpi() -> int:
    raw = os.getenv("BC_SCIENCE_OCR_DPI", str(DEFAULT_OCR_DPI))
    try:
        value = int(raw)
    except ValueError:
        return DEFAULT_OCR_DPI
    return max(96, min(value, 300))


@lru_cache(maxsize=1)
def _default_ocr_language() -> str:
    executable = shutil.which("tesseract")
    if executable:
        try:
            result = subprocess.run(
                [executable, "--list-langs"], check=False, capture_output=True,
                text=True, timeout=5, errors="replace",
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            if result.returncode == 0 and "ita" in result.stdout.splitlines():
                return "ita+eng"
        except (OSError, subprocess.TimeoutExpired):
            pass
    return "eng"


def _ocr_language() -> str:
    return os.getenv("BC_SCIENCE_OCR_LANG") or _default_ocr_language()


def _recover_pdf_pages(path: Path, page_indexes: list[int]) -> dict[int, str]:
    if not page_indexes:
        return {}

    try:
        chunks = pymupdf4llm.to_markdown(
            str(path),
            pages=page_indexes,
            page_chunks=True,
            use_ocr=True,
            ocr_language=_ocr_language(),
            ocr_dpi=_ocr_dpi(),
            show_progress=False,
        )
    except (ImportError, OSError, RuntimeError, TypeError, ValueError):
        # Recovery is deliberately best-effort. Native PyMuPDF and legacy OCR remain available.
        return {}

    if isinstance(chunks, str):
        if len(page_indexes) == 1 and chunks.strip():
            return {page_indexes[0]: chunks.strip()}
        return {}

    recovered: dict[int, str] = {}
    for page_index, chunk in zip(page_indexes, chunks):
        if isinstance(chunk, dict):
            text = str(chunk.get("text", "")).strip()
        else:
            text = str(chunk).strip()
        if text:
            text = re.sub(r"(?i)<br\s*/?>", "\n", text)
            text = re.sub(r"<!--.*?-->", "", text, flags=re.DOTALL)
            recovered[page_index] = text
    return recovered


def _legacy_ocr_page(page: pymupdf.Page) -> str:
    try:
        textpage = page.get_textpage_ocr(
            language=_ocr_language(),
            dpi=_ocr_dpi(),
        )
        return page.get_text(
            "text",
            textpage=textpage,
            sort=True,
        ).strip()
    except (RuntimeError, OSError):
        return ""


def _transcribe_visual_pages(path: Path, page_indexes: list[int]) -> dict[int, str]:
    """Read diagram labels locally when the installed model supports images."""
    if not page_indexes or os.getenv("BC_SCIENCE_VISION_OCR", "1") == "0":
        return {}
    import httpx

    from .ollama_client import OllamaClient

    config = AppConfig.load()
    client = OllamaClient(config.ollama_url)
    try:
        if not client.available():
            return {}
        response = httpx.post(config.ollama_url.rstrip("/") + "/api/show",
            json={"model": config.generation_model}, timeout=10)
        response.raise_for_status()
        if "vision" not in response.json().get("capabilities", []):
            return {}
    except (httpx.HTTPError, ValueError):
        return {}
    recovered = {}
    with pymupdf.open(path) as doc:
        for index in page_indexes:
            image = doc[index].get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes("png")
            try:
                result = client.chat_stream(config.generation_model, [
                    {"role": "system", "content": "Trascrivi testi presenti nelle immagini. Non aggiungere conoscenze esterne e non completare parole illeggibili."},
                    {"role": "user", "content": "Trascrivi solo i testi degli schemi e delle figure, mantenendo separati i riquadri. Ignora intestazioni universitarie e pie di pagina. Se una parola non e leggibile scrivi [illeggibile]. Non interpretare ne spiegare le figure.",
                     "images": [base64.b64encode(image).decode("ascii")]},
                ], num_ctx=config.num_ctx, num_predict=1800, keep_alive=config.keep_alive)
                if result.content.strip() and result.done_reason != "length":
                    recovered[index] = result.content.strip()
            except (RuntimeError, httpx.HTTPError, ValueError):
                continue
    return recovered


def _extract_pdf(path: Path, ocr: bool = True) -> ExtractedDocument:
    with pymupdf.open(path) as doc:
        page_count = len(doc)
        page_texts = [page.get_text("text", sort=True).strip() for page in doc]
        native_texts = list(page_texts)
        image_rows = [page.get_image_info(hashes=True) for page in doc] if ocr else []
        occurrences: dict[bytes, int] = {}
        for rows in image_rows:
            for digest in {row.get("digest", b"") for row in rows}:
                occurrences[digest] = occurrences.get(digest, 0) + 1
        visual_pages = [i for i, rows in enumerate(image_rows) if any(
            occurrences.get(row.get("digest", b""), 0) == 1
            and pymupdf.Rect(row["bbox"]).get_area() / doc[i].rect.get_area() >= 0.12
            for row in rows
        )]

    recovered_pages = 0
    if ocr:
        weak_pages = [
            index
            for index, text in enumerate(page_texts)
            if _needs_pdf_recovery(text)
        ]
        recovered = _recover_pdf_pages(path, weak_pages)
        improved: set[int] = set()

        for page_index, candidate in recovered.items():
            if len(normalize_text(candidate)) > len(normalize_text(page_texts[page_index])):
                page_texts[page_index] = candidate
                improved.add(page_index)
                recovered_pages += 1

        remaining = [index for index in weak_pages if index not in improved]
        if remaining:
            with pymupdf.open(path) as doc:
                for page_index in remaining:
                    candidate = _legacy_ocr_page(doc[page_index])
                    if len(normalize_text(candidate)) > len(
                        normalize_text(page_texts[page_index])
                    ):
                        page_texts[page_index] = candidate
                        recovered_pages += 1

    if ocr:
        visual_texts = _transcribe_visual_pages(path, visual_pages)
        for index in visual_pages:
            if index in visual_texts:
                supplement = page_texts[index] if page_texts[index] != native_texts[index] else ""
                page_texts[index] = native_texts[index] + (
                    "\n\n[Trascrizione visiva automatica della figura: da verificare sull'originale]\n"
                    + visual_texts[index]
                )
                if supplement:
                    page_texts[index] += (
                        "\n\n[Inizio OCR supplementare: possibile rumore, confrontare la figura originale]\n"
                        + supplement + "\n[Fine OCR supplementare]"
                    )
            else:
                page_texts[index] += "\n\n[Figura presente: trascrizione visiva non disponibile; verificare sull'originale.]"

    parts = [
        f"\n--- Pagina {index + 1} ---\n{text}"
        for index, text in enumerate(page_texts)
        if text
    ]
    return ExtractedDocument(
        path,
        normalize_text("\n".join(parts)),
        page_count,
        recovered_pages,
    )


def _extract_docx(path: Path) -> ExtractedDocument:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return ExtractedDocument(path, normalize_text(text))


def extract_document(path: Path, *, ocr: bool = True) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, ocr=ocr)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return ExtractedDocument(
            path,
            normalize_text(path.read_text(encoding="utf-8", errors="replace")),
        )
    raise ValueError(f"Formato non supportato: {path.suffix}")



# Increase when extraction normalization or OCR recovery semantics change.
EXTRACTION_CACHE_VERSION = "extraction-v4-separated-ocr-provenance"


def readable_source(text: str) -> str:
    """Keep every extraction channel but explicitly distinguish uncertain OCR.

    No content is silently deleted: a noisy alternative transcription must not
    masquerade as native document prose or independent confirming evidence.
    The markers remain available to models and the original-page viewer.
    """
    return text.replace(
        "[Inizio OCR supplementare: possibile rumore, confrontare la figura originale]",
        "[Inizio OCR supplementare: trascrizione alternativa NON verificata della stessa figura, "
        "non un secondo fatto. Preferire il testo nativo e la trascrizione visiva; "
        "se discordano segnalare il dubbio senza inventare una riconciliazione.]",
    )


def extraction_signature(path: Path, *, ocr: bool = True) -> str:
    settings = (
        f"|ocr={int(ocr)}|dpi={_ocr_dpi()}"
        f"|lang={_ocr_language()}"
        f"|vision={os.getenv('BC_SCIENCE_VISION_OCR', '1')}:{AppConfig.load().generation_model}"
        if path.suffix.casefold() == ".pdf" else ""
    )
    return EXTRACTION_CACHE_VERSION + settings


def extract_document_cached(
    path: Path,
    db: CacheDB,
    *,
    file_hash: str | None = None,
    ocr: bool = True,
    progress: Callable[[str], None] | None = None,
) -> ExtractedDocument:
    """Cache completed text extraction across indexing and summarization.

    The cache is keyed by original file bytes and OCR settings. It keeps a
    plain-text copy in the existing local SQLite workspace only; original
    documents and embedding rows are untouched.
    """
    path = path.expanduser().resolve()
    digest = file_hash if file_hash is not None else file_sha256(path)
    signature = extraction_signature(path, ocr=ocr)
    cached = db.get_extracted_document(str(path), digest, signature)
    if cached is not None:
        if progress:
            progress(f"    Testo gia estratto (OCR riutilizzato): {path.name}")
        text, pages, ocr_pages = cached
        return ExtractedDocument(path, text, pages, ocr_pages)

    if progress:
        progress(f"    Estrazione del testo (OCR se necessario): {path.name}")
    extracted = extract_document(path, ocr=ocr)
    if extracted.text.strip() and "trascrizione visiva non disponibile" not in extracted.text:
        db.set_extracted_document(
            str(path), digest, signature,
            extracted.text, extracted.pages, extracted.ocr_pages,
        )
    return extracted


def _safe_extract_zip(path: Path) -> Path:
    digest = file_sha256(path)[:16]
    target = app_home() / "imports" / digest
    marker = target / ".ready"
    if marker.exists():
        return target

    temp = target.with_name(target.name + ".tmp")
    if temp.exists():
        shutil.rmtree(temp)
    temp.mkdir(parents=True, exist_ok=True)
    root = temp.resolve()

    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            member = Path(info.filename)
            if member.is_absolute() or ".." in member.parts:
                continue
            destination = (temp / member).resolve()
            if root not in destination.parents and destination != root:
                continue
            if info.is_dir():
                destination.mkdir(parents=True, exist_ok=True)
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(info) as src, destination.open("wb") as dst:
                shutil.copyfileobj(src, dst)

    if target.exists():
        shutil.rmtree(target)
    temp.rename(target)
    marker.touch()
    return target


def iter_source_files(source: Path) -> list[Path]:
    source = source.expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(source)

    if source.is_file() and source.suffix.lower() == ".zip":
        source = _safe_extract_zip(source)

    if source.is_file():
        return (
            [source]
            if source.suffix.lower() in SUPPORTED and not is_generated_artifact(source)
            else []
        )

    return sorted(
        p
        for p in source.rglob("*")
        if (
            p.is_file()
            and p.suffix.lower() in SUPPORTED
            and not is_generated_artifact(p)
        )
    )


def chunk_text(text: str, chunk_chars: int = 5200, overlap: int = 450) -> list[str]:
    if not text.strip():
        return []
    if len(text) <= chunk_chars:
        return [text.strip()]

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > chunk_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            step = max(1, chunk_chars - overlap)
            while start < len(paragraph):
                piece = paragraph[start:start + chunk_chars].strip()
                if piece:
                    chunks.append(piece)
                start += step
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_chars:
            current = candidate
        else:
            chunks.append(current.strip())
            tail = current[-overlap:] if overlap and current else ""
            current = f"{tail}\n\n{paragraph}".strip()

    if current:
        chunks.append(current.strip())
    return chunks
